"""Exportação / importação Excel (.xlsx) — dados mestre + export de documentos."""
from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple
from collections import defaultdict
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.database import new_id, now_iso, next_sequence
from app.repositories import (
    artigos_repo, categorias_repo, clientes_repo, consumiveis_repo,
    encomendas_repo, fornecedores_repo, mao_obra_repo, maquinas_repo, orcamentos_repo,
    ordens_repo, subcategorias_repo,
)
from app.services.numeracao import next_codigo

MAX_ROWS = 5000

# Modos de importação:
# - create: só cria novos (predefinição) — nunca sobrescreve existentes
# - update: actualiza registos com o mesmo código/número (requer confirmação na UI)
IMPORT_MODES = ("create", "update")

# NIFs genéricos / placeholder — nunca usar para match de duplicados
_PLACEHOLDER_NIFS = {
    "", "0", "000000000", "123456789", "999999999", "111111111", "222222222",
}

# Colunas que devem ser texto no Excel (evita 912345678 → 9.12345678E8 ou perda de zeros)
_TEXT_COLUMNS = {
    "codigo", "numero", "encomenda_numero", "artigo_codigo", "orcamento_numero",
    "nif", "contacto", "codigo_postal", "email", "cod_fabricante", "cod_fornecedor",
    "codigo_produto", "plano_contas", "categoria_codigo",
}

# entity_key → módulo RBAC
ENTITY_PERM = {
    "clientes": "clientes",
    "fornecedores": "fornecedores",
    "artigos": "artigos",
    "materiais": "materiais",
    "maquinas": "maquinas",
    "mao_obra": "mao_obra",
    "categorias": "artigos",
    "subcategorias": "artigos",
    "orcamentos": "orcamentos",
    "encomendas": "encomendas",
    "encomenda_linhas": "encomendas",
    "ordens_fabrico": "ordens_fabrico",
}

IMPORTABLE = {
    "clientes", "fornecedores", "artigos", "materiais", "maquinas", "mao_obra",
    "categorias", "subcategorias", "encomendas",
}

EXPORTABLE = set(k for k in ENTITY_PERM.keys() if k not in ("encomenda_linhas", "subcategorias"))

# Folhas e colunas canónicas (ordem = cabeçalho Excel)
SHEETS: Dict[str, Dict[str, Any]] = {
    "clientes": {
        "title": "Clientes",
        "columns": [
            "codigo", "nome", "tipo", "morada", "codigo_postal", "cidade", "pais",
            "contacto", "email", "nif", "notas",
        ],
    },
    "fornecedores": {
        "title": "Fornecedores",
        "columns": [
            "codigo", "nome", "tipo", "morada", "codigo_postal", "cidade", "pais",
            "contacto", "email", "nif", "website", "categoria", "notas",
        ],
    },
    "artigos": {
        "title": "Artigos",
        "columns": [
            "codigo", "nome", "descricao", "unidade", "tipo_artigo", "produzido",
            "categoria", "subcategoria",
            "custo_artigo", "margem", "comissao_pct",
            "fabricante", "cod_fabricante", "fornecedor", "cod_fornecedor",
            "website", "ficha_produto", "plano_contas", "codigo_produto",
            "comprimento_mm", "largura_mm", "espessura_mm", "peso_kg",
            "comprimento_unidade", "largura_unidade", "espessura_unidade", "peso_unidade",
            "responsavel", "qtd_uni", "qtd_stock", "nivel_reabastecimento", "qtd_ultima_compra",
            "diversos", "ativo",
        ],
    },
    "materiais": {
        "title": "Materiais",
        "columns": ["codigo", "nome", "unidade", "custo_unitario"],
    },
    "maquinas": {
        "title": "Maquinas",
        "columns": ["codigo", "nome", "custo_amortizacao_hora", "custo_energia_hora"],
    },
    "mao_obra": {
        "title": "Mao_de_obra",
        "columns": ["codigo", "nome", "custo_hora", "responsavel_personalizacoes"],
    },
    "categorias": {
        "title": "Categorias",
        "columns": ["codigo", "nome"],
    },
    "subcategorias": {
        "title": "Subcategorias",
        "columns": ["codigo", "nome", "categoria_codigo", "categoria"],
    },
    "orcamentos": {
        "title": "Orcamentos",
        "columns": [
            "numero", "cliente", "descricao", "numero_encomenda", "data", "validade",
            "status", "margem", "total", "notas",
        ],
    },
    "encomendas": {
        "title": "Encomendas",
        "columns": [
            "numero", "cliente", "descricao", "data", "prazo_entrega", "estado",
            "valor_total", "valor_pago", "desconto_total", "desconto_total_tipo", "envio",
            "orcamento_numero", "notas",
        ],
    },
    "encomenda_linhas": {
        "title": "Encomenda_linhas",
        "columns": [
            "encomenda_numero", "artigo_codigo", "artigo_nome", "quantidade",
            "preco_unit", "desconto", "desconto_tipo", "personalizacoes",
        ],
    },
    "ordens_fabrico": {
        "title": "Ordens_fabrico",
        "columns": [
            "numero", "cliente", "descricao", "data", "status", "prioritaria",
            "responsavel_nome", "encomenda_numero", "orcamento_numero", "notas",
        ],
    },
}

# Ao pedir "categorias" / "encomendas" no hub, exporta folhas filhas
EXPORT_EXPAND = {
    "categorias": ["categorias", "subcategorias"],
    "encomendas": ["encomendas", "encomenda_linhas"],
}


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "sim" if v else "nao"
    return v


def _fix_mojibake(s: str) -> str:
    """Corrige texto UTF-8 mal interpretado (ex.: GuimarÃ£es → Guimarães, CÃƒO → CÃO)."""
    if not s:
        return s
    cur = s
    for _ in range(3):
        if "Ã" not in cur and "Â" not in cur and "â€" not in cur and "ƒ" not in cur:
            break
        nxt = None
        # cp1252 primeiro: trata ƒ (U+0192) em CÃƒO / Ãƒ
        for enc in ("cp1252", "latin-1"):
            try:
                cand = cur.encode(enc).decode("utf-8")
                if cand and cand != cur:
                    nxt = cand
                    break
            except (UnicodeEncodeError, UnicodeDecodeError):
                continue
        if not nxt:
            break
        cur = nxt
    # Restos típicos: Á exportado mal como Ã antes de consoante (GRÃFICA → GRÁFICA)
    # Não mexe em ÃO/ÃE/ÃA (CÃO, etc.)
    cur = re.sub(r"Ã(?=[BCDFGHJKLMNPQRSTVWXYZbcdfghjklmnpqrstvwxyz])", "Á", cur)
    cur = cur.replace("Ã§", "ç").replace("Ã‡", "Ç")
    return cur


def _norm_str(v: Any) -> str:
    """Normaliza qualquer célula para texto limpo (códigos, nomes, contactos)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "sim" if v else "nao"
    if isinstance(v, int) and not isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return str(v).rstrip("0").rstrip(".") if "." in str(v) else str(v)
    s = str(v).strip()
    if not s or s.lower() in ("none", "null", "nan"):
        return ""
    # Excel por vezes guarda "123456789.0"
    if re.fullmatch(r"-?\d+\.0+", s):
        s = s.split(".", 1)[0]
    return _fix_mojibake(s)


def _norm_code(v: Any) -> str:
    return _norm_str(v).strip()


def _parse_bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    s = _norm_str(v).lower()
    return s in ("1", "true", "sim", "yes", "s", "y")


def _parse_float(v: Any, default: float = 0.0) -> float:
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return float(v)
        except (TypeError, ValueError):
            return default
    s = _norm_str(v)
    if not s:
        return default
    # Europeu: 1.234,56 ou 12,5
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s and s.count(",") == 1:
        s = s.replace(",", ".")
    s = re.sub(r"[^\d.\-]", "", s)
    try:
        return float(s)
    except (TypeError, ValueError):
        return default


def _parse_date(v: Any) -> Optional[str]:
    if v is None or v == "":
        return None
    if hasattr(v, "strftime"):
        try:
            return v.strftime("%Y-%m-%d")
        except Exception:
            return None
    s = _norm_str(v)
    if not s:
        return None
    # dd/mm/yyyy or dd-mm-yyyy
    for sep in ("/", "-"):
        parts = s.split(sep)
        if len(parts) == 3 and len(parts[2]) == 4:
            d, m, y = parts[0], parts[1], parts[2]
            if len(d) <= 2 and len(m) <= 2:
                try:
                    return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
                except ValueError:
                    pass
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    return s[:10] if len(s) >= 8 else s


def _is_placeholder_nif(nif: str) -> bool:
    digits = re.sub(r"\D", "", nif or "")
    return digits in _PLACEHOLDER_NIFS or len(digits) < 5


def _norm_nif(v: Any) -> Optional[str]:
    """NIF real ou None (vazio / placeholder tipo 123456789 → null)."""
    s = _norm_code(v)
    if not s or _is_placeholder_nif(s):
        return None
    return re.sub(r"\D", "", s) or None


def _row_wants_update(row: Dict[str, Any], mode: str) -> bool:
    """Actualizar só em mode=update ou se a coluna atualizar=sim."""
    flag = _norm_str(row.get("atualizar") or row.get("actualizar") or "").lower()
    if flag in ("1", "sim", "yes", "s", "y", "true", "update", "actualizar", "atualizar"):
        return True
    return mode == "update"


def _is_example_row(row: Dict[str, Any]) -> bool:
    for key in ("codigo", "numero", "encomenda_numero", "nome", "cliente"):
        v = _norm_str(row.get(key) or "")
        if v.startswith("#") or v.upper().startswith("EXEMPLO"):
            return True
    return False


def _row_dict(headers: List[str], values: tuple) -> Dict[str, Any]:
    out = {}
    for i, h in enumerate(headers):
        key = (h or "").strip().lower()
        if not key:
            continue
        raw = values[i] if i < len(values) else None
        if key in _TEXT_COLUMNS or key in (
            "nome", "cliente", "descricao", "morada", "cidade", "pais",
            "tipo", "estado", "notas", "artigo_nome", "personalizacoes",
            "categoria", "subcategoria", "fornecedor", "fabricante",
            "atualizar", "actualizar",
        ):
            out[key] = _norm_str(raw)
        else:
            out[key] = raw
    return out


async def _load_docs(entity: str, ids: Optional[List[str]] = None) -> List[dict]:
    repo_map = {
        "clientes": clientes_repo,
        "fornecedores": fornecedores_repo,
        "artigos": artigos_repo,
        "materiais": consumiveis_repo,
        "maquinas": maquinas_repo,
        "mao_obra": mao_obra_repo,
        "categorias": categorias_repo,
        "subcategorias": subcategorias_repo,
        "orcamentos": orcamentos_repo,
        "encomendas": encomendas_repo,
        "ordens_fabrico": ordens_repo,
    }
    repo = repo_map[entity]
    query: dict = {}
    if ids is not None:
        if not ids:
            return []
        query = {"id": {"$in": ids}}
    return await repo.find(query, sort=("created_at", 1), limit=MAX_ROWS)


def _doc_to_row(entity: str, doc: dict) -> List[Any]:
    cols = SHEETS[entity]["columns"]
    if entity == "clientes":
        m = {
            "codigo": doc.get("codigo"),
            "nome": doc.get("nome"),
            "tipo": doc.get("tipo") or ("empresa" if doc.get("nif") else "particular"),
            "morada": doc.get("morada"),
            "codigo_postal": doc.get("codigo_postal"),
            "cidade": doc.get("cidade"),
            "pais": doc.get("pais"),
            "contacto": doc.get("contacto"),
            "email": doc.get("email"),
            "nif": doc.get("nif"),
            "notas": doc.get("notas"),
        }
    elif entity == "fornecedores":
        m = {
            "codigo": doc.get("codigo"),
            "nome": doc.get("nome"),
            "tipo": doc.get("tipo") or ("empresa" if doc.get("nif") else "particular"),
            "morada": doc.get("morada"),
            "codigo_postal": doc.get("codigo_postal"),
            "cidade": doc.get("cidade"),
            "pais": doc.get("pais"),
            "contacto": doc.get("contacto"),
            "email": doc.get("email"),
            "nif": doc.get("nif"),
            "website": doc.get("website"),
            "categoria": doc.get("categoria"),
            "notas": doc.get("notas"),
        }
    elif entity == "artigos":
        m = {
            "codigo": doc.get("codigo"),
            "nome": doc.get("nome"),
            "descricao": doc.get("descricao"),
            "unidade": doc.get("unidade") or "un",
            "tipo_artigo": doc.get("tipo_artigo") or "ativo",
            "produzido": doc.get("produzido"),
            "categoria": doc.get("categoria_nome"),
            "subcategoria": doc.get("subcategoria_nome"),
            "custo_artigo": doc.get("custo_artigo"),
            "margem": doc.get("margem"),
            "comissao_pct": doc.get("comissao_pct"),
            "fabricante": doc.get("fabricante"),
            "cod_fabricante": doc.get("cod_fabricante"),
            "fornecedor": doc.get("fornecedor_nome"),
            "cod_fornecedor": doc.get("cod_fornecedor"),
            "website": doc.get("website"),
            "ficha_produto": doc.get("ficha_produto"),
            "plano_contas": doc.get("plano_contas"),
            "codigo_produto": doc.get("codigo_produto"),
            "comprimento_mm": doc.get("comprimento_mm"),
            "largura_mm": doc.get("largura_mm"),
            "espessura_mm": doc.get("espessura_mm"),
            "peso_kg": doc.get("peso_kg"),
            "comprimento_unidade": doc.get("comprimento_unidade") or "mm",
            "largura_unidade": doc.get("largura_unidade") or "mm",
            "espessura_unidade": doc.get("espessura_unidade") or "mm",
            "peso_unidade": doc.get("peso_unidade") or "kg",
            "responsavel": doc.get("responsavel"),
            "qtd_uni": doc.get("qtd_uni"),
            "qtd_stock": doc.get("qtd_stock"),
            "nivel_reabastecimento": doc.get("nivel_reabastecimento"),
            "qtd_ultima_compra": doc.get("qtd_ultima_compra"),
            "diversos": doc.get("diversos"),
            "ativo": doc.get("ativo") if doc.get("ativo") is not None else True,
        }
    elif entity == "materiais":
        m = {
            "codigo": doc.get("codigo"),
            "nome": doc.get("nome"),
            "unidade": doc.get("unidade") or "un",
            "custo_unitario": doc.get("custo_unitario"),
        }
    elif entity == "maquinas":
        m = {
            "codigo": doc.get("codigo"),
            "nome": doc.get("nome"),
            "custo_amortizacao_hora": doc.get("custo_amortizacao_hora"),
            "custo_energia_hora": doc.get("custo_energia_hora"),
        }
    elif entity == "mao_obra":
        m = {
            "codigo": doc.get("codigo"),
            "nome": doc.get("nome"),
            "custo_hora": doc.get("custo_hora"),
            "responsavel_personalizacoes": doc.get("responsavel_personalizacoes"),
        }
    elif entity == "categorias":
        m = {"codigo": doc.get("codigo"), "nome": doc.get("nome")}
    elif entity == "subcategorias":
        m = {
            "codigo": doc.get("codigo"),
            "nome": doc.get("nome"),
            "categoria_codigo": "",  # preenchido abaixo se necessário
            "categoria": doc.get("categoria_nome"),
        }
    elif entity == "orcamentos":
        m = {
            "numero": doc.get("numero"),
            "cliente": doc.get("cliente"),
            "descricao": doc.get("descricao"),
            "numero_encomenda": doc.get("numero_encomenda"),
            "data": doc.get("data"),
            "validade": doc.get("validade"),
            "status": doc.get("status"),
            "margem": doc.get("margem"),
            "total": doc.get("total"),
            "notas": doc.get("notas"),
        }
    elif entity == "encomendas":
        m = {
            "numero": doc.get("numero"),
            "cliente": doc.get("cliente"),
            "descricao": doc.get("descricao"),
            "data": doc.get("data"),
            "prazo_entrega": doc.get("prazo_entrega"),
            "estado": doc.get("estado"),
            "valor_total": doc.get("valor_total"),
            "valor_pago": doc.get("valor_pago"),
            "desconto_total": doc.get("desconto_total"),
            "desconto_total_tipo": doc.get("desconto_total_tipo") or "pct",
            "envio": doc.get("envio"),
            "orcamento_numero": doc.get("orcamento_numero"),
            "notas": doc.get("notas"),
        }
    elif entity == "ordens_fabrico":
        m = {
            "numero": doc.get("numero"),
            "cliente": doc.get("cliente"),
            "descricao": doc.get("descricao"),
            "data": doc.get("data"),
            "status": doc.get("status"),
            "prioritaria": doc.get("prioritaria"),
            "responsavel_nome": doc.get("responsavel_nome"),
            "encomenda_numero": doc.get("encomenda_numero"),
            "orcamento_numero": doc.get("orcamento_numero"),
            "notas": doc.get("notas"),
        }
    else:
        m = {}
    return [_cell(m.get(c)) for c in cols]


def _pers_to_cell(pers: Any) -> str:
    parts = []
    for p in pers or []:
        nome = (p.get("nome") or "").strip()
        if not nome:
            continue
        parts.append(f"{nome}={_parse_float(p.get('valor'))}")
    return ";".join(parts)


def _pers_from_cell(raw: Any) -> List[dict]:
    s = str(raw or "").strip()
    if not s:
        return []
    out = []
    for part in s.split(";"):
        part = part.strip()
        if not part:
            continue
        if "=" in part:
            n, v = part.split("=", 1)
        elif ":" in part:
            n, v = part.split(":", 1)
        else:
            n, v = part, "0"
        out.append({
            "id": new_id(),
            "nome": n.strip(),
            "valor": _parse_float(v),
            "tempo": 0,
        })
    return out


def _doc_to_row_enc_linha(enc: dict, linha: dict) -> List[Any]:
    cols = SHEETS["encomenda_linhas"]["columns"]
    m = {
        "encomenda_numero": enc.get("numero"),
        "artigo_codigo": linha.get("_artigo_codigo") or "",
        "artigo_nome": linha.get("artigo_nome"),
        "quantidade": linha.get("quantidade"),
        "preco_unit": linha.get("preco_unit"),
        "desconto": linha.get("desconto"),
        "desconto_tipo": linha.get("desconto_tipo") or "pct",
        "personalizacoes": _pers_to_cell(linha.get("personalizacoes")),
    }
    return [_cell(m.get(c)) for c in cols]


async def _enrich_subcategorias(docs: List[dict]) -> List[dict]:
    """Acrescenta categoria_codigo via lookup."""
    cats = {c["id"]: c for c in await categorias_repo.find(limit=MAX_ROWS)}
    for d in docs:
        cat = cats.get(d.get("categoria_id") or "")
        d["_categoria_codigo"] = (cat or {}).get("codigo") or ""
        if not d.get("categoria_nome") and cat:
            d["categoria_nome"] = cat.get("nome") or ""
    return docs


def _doc_to_row_sub(doc: dict) -> List[Any]:
    cols = SHEETS["subcategorias"]["columns"]
    m = {
        "codigo": doc.get("codigo"),
        "nome": doc.get("nome"),
        "categoria_codigo": doc.get("_categoria_codigo"),
        "categoria": doc.get("categoria_nome"),
    }
    return [_cell(m.get(c)) for c in cols]


def _style_header(ws, n_cols: int):
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="111111")
    for col in range(1, n_cols + 1):
        cell = ws.cell(1, col)
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}1"


def _format_text_columns(ws, columns: List[str], *, prefill_rows: int = 80) -> None:
    """Força formato Texto nas colunas sensíveis para o Excel não estragar códigos/NIF/telefone."""
    for idx, name in enumerate(columns, start=1):
        if name not in _TEXT_COLUMNS:
            continue
        letter = get_column_letter(idx)
        for r in range(2, prefill_rows + 2):
            ws.cell(r, idx).number_format = "@"
        ws.column_dimensions[letter].width = max(14, min(28, len(name) + 6))


def _add_sheet_validations(ws, entity: str, columns: List[str]) -> None:
    """Dropdowns para reduzir erros de preenchimento."""
    col_index = {c: i + 1 for i, c in enumerate(columns)}
    specs = {
        "clientes": {"tipo": '"empresa,particular"'},
        "fornecedores": {"tipo": '"empresa,particular"'},
        "encomendas": {"estado": '"aberta,em_producao,concluida,cancelada"', "desconto_total_tipo": '"pct,eur"'},
        "encomenda_linhas": {"desconto_tipo": '"pct,eur"'},
        "artigos": {
            "tipo_artigo": '"ativo,consumivel,servico,nao_utilizado,inativo"',
            "produzido": '"sim,nao"',
            "ativo": '"sim,nao"',
        },
    }.get(entity, {})
    for field, formula in specs.items():
        idx = col_index.get(field)
        if not idx:
            continue
        letter = get_column_letter(idx)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Valor inválido"
        dv.errorTitle = "Velocely"
        dv.prompt = "Escolha da lista"
        dv.promptTitle = field
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}5001")


def _add_sheet(wb: Workbook, entity: str, rows: List[List[Any]], *, first: bool) -> None:
    meta = SHEETS[entity]
    title = meta["title"][:31]
    if first and wb.active.title == "Sheet":
        ws = wb.active
        ws.title = title
    else:
        ws = wb.create_sheet(title)
    cols = meta["columns"]
    ws.append(cols)
    _style_header(ws, len(cols))
    for row in rows:
        # Garantir códigos/NIF como texto mesmo quando vêm numéricos
        safe = []
        for i, val in enumerate(row):
            col = cols[i] if i < len(cols) else ""
            if col in _TEXT_COLUMNS and val is not None and val != "":
                safe.append(_norm_str(val))
            else:
                safe.append(val)
        ws.append(safe)
        r = ws.max_row
        for i, col in enumerate(cols, start=1):
            if col in _TEXT_COLUMNS:
                ws.cell(r, i).number_format = "@"
                # Prefixo ' implícito via number_format; valor já é str
    _format_text_columns(ws, cols, prefill_rows=max(80, min(len(rows) + 20, 120)))
    _add_sheet_validations(ws, entity, cols)


def _instructions_sheet(wb: Workbook, entities: List[str]) -> None:
    ws = wb.create_sheet("Instrucoes", 0)
    lines = [
        ["Velocely — Importação / Exportação Excel"],
        [""],
        ["Como preencher (para evitar erros):"],
        ["1. Não altere os nomes das colunas na 1.ª linha."],
        ["2. Códigos, NIF, telefone e código postal: deixe como TEXTO (o template já formata assim)."],
        ["3. Por defeito a importação SÓ CRIA registos novos — nunca substitui os que já existem."],
        ["4. Para actualizar um existente: use o modo «Actualizar» na app OU coloque atualizar=sim na linha."],
        ["6. NIF: pode ficar vazio (grava null). Valores falsos como 123456789 são ignorados (null)."],
        ["7. Identidade: clientes/artigos pelo campo codigo; encomendas pelo campo numero."],
        ["8. Encomendas: folha Encomendas (1 linha por encomenda) + Encomenda_linhas (várias linhas por numero)."],
        ["9. Se a folha Encomendas tiver o mesmo numero repetido (1× por artigo), o sistema usa só a 1.ª ocorrência."],
        ["10. Linhas ligam-se por encomenda_numero = numero da encomenda. Artigos por artigo_codigo."],
        ["11. Datas: AAAA-MM-DD ou DD/MM/AAAA. Máximo 5000 linhas por folha."],
        ["12. Linhas de exemplo começadas por # ou EXEMPLO são ignoradas."],
        [""],
        ["Folhas neste ficheiro:"],
    ]
    for e in entities:
        lines.append([SHEETS[e]["title"], ", ".join(SHEETS[e]["columns"])])
    for row in lines:
        ws.append(row)
    ws.column_dimensions["A"].width = 100
    ws.column_dimensions["B"].width = 80
    ws["A1"].font = Font(bold=True, size=14)


async def build_export_xlsx(
    entities: List[str],
    ids_by_entity: Optional[Dict[str, List[str]]] = None,
) -> bytes:
    """Gera .xlsx com uma folha por entidade (expandindo categorias / encomendas)."""
    expanded: List[str] = []
    for e in entities:
        e = (e or "").strip()
        if not e:
            continue
        if e not in SHEETS and e not in EXPORT_EXPAND:
            continue
        for x in EXPORT_EXPAND.get(e, [e]):
            if x in SHEETS and x not in expanded:
                expanded.append(x)
    if not expanded:
        raise ValueError("Seleccione pelo menos uma entidade válida")

    wb = Workbook()
    first = True
    for entity in expanded:
        id_list = None
        if ids_by_entity is not None:
            if entity == "encomenda_linhas":
                id_list = ids_by_entity.get("encomendas") or ids_by_entity.get("encomenda_linhas")
            elif entity in ids_by_entity:
                id_list = ids_by_entity[entity]

        if entity == "encomenda_linhas":
            encs = await _load_docs("encomendas", id_list)
            art_ids = {a.get("artigo_id") for e in encs for a in (e.get("artigos") or []) if a.get("artigo_id")}
            arts = await artigos_repo.find({"id": {"$in": list(art_ids)}}, limit=MAX_ROWS) if art_ids else []
            art_cod = {a["id"]: a.get("codigo") or "" for a in arts}
            rows = []
            for enc in encs:
                for linha in enc.get("artigos") or []:
                    linha = {**linha, "_artigo_codigo": art_cod.get(linha.get("artigo_id") or "", "")}
                    rows.append(_doc_to_row_enc_linha(enc, linha))
            _add_sheet(wb, entity, rows, first=first)
            first = False
            continue

        docs = await _load_docs(entity, id_list)
        if entity == "subcategorias":
            docs = await _enrich_subcategorias(docs)
            rows = [_doc_to_row_sub(d) for d in docs]
        else:
            rows = [_doc_to_row(entity, d) for d in docs]
        _add_sheet(wb, entity, rows, first=first)
        first = False

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template_xlsx(entity: str) -> bytes:
    if entity not in IMPORTABLE:
        raise ValueError("Entidade não suportada para importação")
    keys = EXPORT_EXPAND.get(entity, [entity])
    wb = Workbook()
    first = True
    for e in keys:
        _add_sheet(wb, e, [], first=first)
        first = False
    _instructions_sheet(wb, keys)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _resolve_categoria(nome_or_codigo: str) -> Optional[dict]:
    s = (nome_or_codigo or "").strip()
    if not s:
        return None
    cat = await categorias_repo.find_one({"codigo": s})
    if cat:
        return cat
    return await categorias_repo.find_one({"nome": s})


async def _resolve_subcategoria(nome: str, categoria_id: Optional[str]) -> Optional[dict]:
    s = (nome or "").strip()
    if not s:
        return None
    q: dict = {"nome": s}
    if categoria_id:
        q["categoria_id"] = categoria_id
    return await subcategorias_repo.find_one(q)


async def import_rows(
    entity: str,
    file_bytes: bytes,
    *,
    dry_run: bool = True,
    mode: str = "create",
) -> Dict[str, Any]:
    """Valida (e opcionalmente grava) linhas de um .xlsx para uma entidade mestre.

    mode=create (predefinição): nunca sobrescreve existentes; reporta skip.
    mode=update: actualiza pelo código/número (ou se a linha tiver atualizar=sim).
    """
    if entity not in IMPORTABLE:
        raise ValueError("Entidade não suportada para importação")
    mode = (mode or "create").strip().lower()
    if mode not in IMPORT_MODES:
        raise ValueError("mode deve ser 'create' ou 'update'")

    targets = EXPORT_EXPAND.get(entity, [entity])

    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    errors: List[dict] = []
    preview: List[dict] = []
    created = updated = skipped = 0
    # Cache de lookups por ficheiro (performance + consistência)
    seen_codes: Dict[str, set] = defaultdict(set)
    # Encomendas criadas neste import (para ligar linhas no mesmo ficheiro)
    created_enc_nums: set = set()

    for target in targets:
        sheet_title = SHEETS[target]["title"]
        ws = None
        for name in wb.sheetnames:
            if name.lower() == sheet_title.lower() or name.lower() == target.lower():
                ws = wb[name]
                break
        if ws is None and len(targets) == 1 and wb.sheetnames:
            for name in wb.sheetnames:
                if name.lower() not in ("instrucoes", "instruções", "instructions"):
                    ws = wb[name]
                    break
        if ws is None:
            if target == "encomenda_linhas":
                continue
            errors.append({"linha": 0, "folha": sheet_title, "erro": f"Folha '{sheet_title}' não encontrada"})
            continue

        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            errors.append({"linha": 0, "folha": sheet_title, "erro": "Folha vazia"})
            continue

        headers = [str(h).strip().lower() if h is not None else "" for h in header_row]
        expected = SHEETS[target]["columns"]
        required = {
            "artigos": ["nome"],
            "encomendas": ["cliente"],
            "encomenda_linhas": ["encomenda_numero"],
            "clientes": ["nome"],
            "fornecedores": ["nome"],
        }.get(target, expected)
        missing = [c for c in required if c not in headers]
        if missing:
            errors.append({
                "linha": 1,
                "folha": sheet_title,
                "erro": f"Colunas em falta: {', '.join(missing)}",
            })
            continue

        # Cache de existentes (evita N queries e acelera ficheiros grandes)
        existing_by_codigo: Dict[str, dict] = {}
        if target == "clientes":
            for d in await clientes_repo.find(limit=MAX_ROWS):
                c = _norm_code(d.get("codigo"))
                if c:
                    existing_by_codigo[c] = d
        elif target == "fornecedores":
            for d in await fornecedores_repo.find(limit=MAX_ROWS):
                c = _norm_code(d.get("codigo"))
                if c:
                    existing_by_codigo[c] = d
        elif target == "artigos":
            for d in await artigos_repo.find(limit=MAX_ROWS):
                c = _norm_code(d.get("codigo"))
                if c:
                    existing_by_codigo[c] = d
        elif target == "encomendas":
            for d in await encomendas_repo.find(limit=MAX_ROWS):
                c = _norm_code(d.get("numero"))
                if c:
                    existing_by_codigo[c] = d

        if target == "encomenda_linhas":
            lines_by_num: Dict[str, List[dict]] = defaultdict(list)
            line_no = 1
            count = 0
            for values in rows_iter:
                line_no += 1
                if not values or all(v is None or str(v).strip() == "" for v in values):
                    continue
                count += 1
                if count > MAX_ROWS:
                    errors.append({"linha": line_no, "folha": sheet_title, "erro": f"Limite de {MAX_ROWS} linhas excedido"})
                    break
                row = _row_dict(headers, values)
                if _is_example_row(row):
                    continue
                num = _norm_code(row.get("encomenda_numero"))
                if not num:
                    errors.append({"linha": line_no, "folha": sheet_title, "erro": "encomenda_numero obrigatório"})
                    continue
                try:
                    linha = await _parse_encomenda_linha(row)
                except Exception as ex:
                    errors.append({"linha": line_no, "folha": sheet_title, "erro": str(ex)})
                    continue
                if isinstance(linha, str):
                    errors.append({"linha": line_no, "folha": sheet_title, "erro": linha})
                    continue
                lines_by_num[num].append(linha)

            for num, linhas in lines_by_num.items():
                try:
                    action, summary = await _apply_encomenda_linhas(
                        num, linhas, dry_run=dry_run, mode=mode,
                        known_new=num in created_enc_nums,
                    )
                except Exception as ex:
                    errors.append({"linha": 0, "folha": sheet_title, "erro": f"{num}: {ex}"})
                    continue
                if action == "error":
                    errors.append({"linha": 0, "folha": sheet_title, "erro": summary if isinstance(summary, str) else str(summary)})
                    continue
                if action == "create":
                    created += 1
                elif action == "update":
                    updated += 1
                elif action == "skip":
                    skipped += 1
                preview.append({"linha": 0, "folha": sheet_title, "acao": action, **(summary if isinstance(summary, dict) else {"nome": summary})})
            continue

        line_no = 1
        count = 0
        for values in rows_iter:
            line_no += 1
            if not values or all(v is None or str(v).strip() == "" for v in values):
                continue
            count += 1
            if count > MAX_ROWS:
                errors.append({"linha": line_no, "folha": sheet_title, "erro": f"Limite de {MAX_ROWS} linhas excedido"})
                break

            row = _row_dict(headers, values)
            if _is_example_row(row):
                continue

            # Dedup dentro do ficheiro (encomendas repetidas 1× por linha de artigo)
            id_key = _norm_code(row.get("numero") if target == "encomendas" else row.get("codigo"))
            if target == "encomendas" and id_key:
                if id_key in seen_codes[target]:
                    skipped += 1
                    preview.append({
                        "linha": line_no, "folha": sheet_title, "acao": "skip",
                        "codigo": id_key, "nome": _norm_str(row.get("cliente")),
                        "motivo": "numero duplicado neste ficheiro (mantida a 1.ª ocorrência)",
                    })
                    continue
                seen_codes[target].add(id_key)
            elif id_key and target in ("clientes", "fornecedores", "artigos", "materiais", "maquinas", "mao_obra", "categorias"):
                if id_key in seen_codes[target]:
                    errors.append({
                        "linha": line_no, "folha": sheet_title,
                        "erro": f"Código duplicado neste ficheiro: {id_key}",
                    })
                    continue
                seen_codes[target].add(id_key)

            try:
                action, summary = await _upsert_row(
                    target, row, dry_run=dry_run, mode=mode,
                    existing_by_codigo=existing_by_codigo,
                )
            except Exception as ex:
                errors.append({"linha": line_no, "folha": sheet_title, "erro": str(ex)})
                continue

            if action == "error":
                errors.append({"linha": line_no, "folha": sheet_title, "erro": summary})
                continue
            if action == "create":
                created += 1
                if target == "encomendas" and isinstance(summary, dict):
                    created_enc_nums.add(_norm_code(summary.get("codigo")))
                if isinstance(summary, dict) and summary.get("codigo") and summary["codigo"] != "(novo)":
                    existing_by_codigo[_norm_code(summary["codigo"])] = {
                        "codigo": summary.get("codigo"),
                        "numero": summary.get("codigo"),
                        "nome": summary.get("nome"),
                        "_pending": True,
                    }
            elif action == "update":
                updated += 1
            elif action == "skip":
                skipped += 1
            preview.append({"linha": line_no, "folha": sheet_title, "acao": action, **summary})

    return {
        "ok": len(errors) == 0,
        "dry_run": dry_run,
        "mode": mode,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "preview": preview[:200],
        "preview_total": len(preview),
    }


async def _upsert_row(
    entity: str,
    row: Dict[str, Any],
    *,
    dry_run: bool,
    mode: str = "create",
    existing_by_codigo: Optional[Dict[str, dict]] = None,
) -> Tuple[str, Any]:
    nome = _norm_str(row.get("nome"))
    codigo = _norm_code(row.get("codigo"))
    allow_update = _row_wants_update(row, mode)
    cache = existing_by_codigo if existing_by_codigo is not None else {}

    if entity == "clientes":
        if not nome:
            return "error", "Nome obrigatório"
        patch = {
            "nome": nome,
            "morada": _norm_str(row.get("morada")),
            "codigo_postal": _norm_code(row.get("codigo_postal")),
            "cidade": _norm_str(row.get("cidade")),
            "pais": _norm_str(row.get("pais")) or "Portugal",
            "contacto": _norm_code(row.get("contacto")),
            "email": _norm_str(row.get("email")).lower(),
            "nif": _norm_nif(row.get("nif")),
            "notas": _norm_str(row.get("notas")),
        }
        tipo_raw = _norm_str(row.get("tipo")).lower().replace(" ", "_")
        if tipo_raw in ("cliente_final", "cliente final"):
            tipo_raw = "particular"
        if tipo_raw in ("empresa", "particular"):
            patch["tipo"] = tipo_raw
        else:
            patch["tipo"] = "empresa" if patch["nif"] else "particular"
        if patch["tipo"] == "empresa":
            faltam = [k for k, v in (
                ("NIF", patch["nif"]),
                ("morada", patch["morada"]),
                ("código postal", patch["codigo_postal"]),
                ("cidade", patch["cidade"]),
                ("contacto", patch["contacto"]),
                ("email", patch["email"]),
            ) if not v]
            if faltam:
                return "error", "Para empresas são obrigatórios: " + ", ".join(faltam)

        existing = cache.get(codigo) if codigo else None
        if existing is None and codigo:
            existing = await clientes_repo.find_one({"codigo": codigo})
        # NIF real noutro código → erro (nunca merge silencioso). Vazio/placeholder → null.
        if not existing and patch["nif"]:
            by_nif = await clientes_repo.find_one({"nif": patch["nif"]})
            if by_nif and _norm_code(by_nif.get("codigo")) != codigo:
                return "error", (
                    f"NIF {patch['nif']} já usado por {_norm_code(by_nif.get('codigo')) or 'outro cliente'} "
                    f"({by_nif.get('nome')}). Não é feito merge automático."
                )

        if existing and not existing.get("_pending"):
            if not allow_update:
                return "skip", {
                    "codigo": existing.get("codigo") or codigo,
                    "nome": existing.get("nome") or nome,
                    "motivo": "já existe — use modo Actualizar ou coluna atualizar=sim",
                }
            if not dry_run:
                await clientes_repo.update(existing["id"], patch)
            return "update", {"codigo": existing.get("codigo") or codigo, "nome": nome}
        if not dry_run:
            doc = {
                "id": new_id(),
                "codigo": codigo or await next_codigo("cliente"),
                **patch,
                "created_at": now_iso(),
            }
            await clientes_repo.insert(doc)
            return "create", {"codigo": doc["codigo"], "nome": nome}
        return "create", {"codigo": codigo or "(novo)", "nome": nome}

    if entity == "fornecedores":
        if not nome:
            return "error", "Nome obrigatório"
        patch = {
            "nome": nome,
            "morada": _norm_str(row.get("morada")),
            "codigo_postal": _norm_code(row.get("codigo_postal")),
            "cidade": _norm_str(row.get("cidade")),
            "pais": _norm_str(row.get("pais")) or "Portugal",
            "contacto": _norm_code(row.get("contacto")),
            "email": _norm_str(row.get("email")).lower(),
            "nif": _norm_nif(row.get("nif")),
            "website": _norm_str(row.get("website")),
            "categoria": _norm_str(row.get("categoria")),
            "notas": _norm_str(row.get("notas")),
        }
        tipo_raw = _norm_str(row.get("tipo")).lower().replace(" ", "_")
        if tipo_raw in ("empresa", "particular"):
            patch["tipo"] = tipo_raw
        else:
            patch["tipo"] = "empresa" if patch["nif"] else "particular"
        if patch["tipo"] == "empresa":
            faltam = [k for k, v in (
                ("NIF", patch["nif"]),
                ("morada", patch["morada"]),
                ("código postal", patch["codigo_postal"]),
                ("cidade", patch["cidade"]),
                ("contacto", patch["contacto"]),
                ("email", patch["email"]),
            ) if not v]
            if faltam:
                return "error", "Para empresas são obrigatórios: " + ", ".join(faltam)
        existing = await fornecedores_repo.find_one({"codigo": codigo}) if codigo else None
        if not existing and patch["nif"]:
            by_nif = await fornecedores_repo.find_one({"nif": patch["nif"]})
            if by_nif and _norm_code(by_nif.get("codigo")) != codigo:
                return "error", (
                    f"NIF {patch['nif']} já usado por {_norm_code(by_nif.get('codigo')) or 'outro fornecedor'} "
                    f"({by_nif.get('nome')})."
                )
        if existing:
            if not allow_update:
                return "skip", {
                    "codigo": existing.get("codigo") or codigo,
                    "nome": existing.get("nome") or nome,
                    "motivo": "já existe — use modo Actualizar ou coluna atualizar=sim",
                }
            if not dry_run:
                await fornecedores_repo.update(existing["id"], patch)
            return "update", {"codigo": existing.get("codigo") or codigo, "nome": nome}
        if not dry_run:
            doc = {
                "id": new_id(),
                "codigo": codigo or await next_codigo("fornecedor"),
                **patch,
                "created_at": now_iso(),
            }
            await fornecedores_repo.insert(doc)
            return "create", {"codigo": doc["codigo"], "nome": nome}
        return "create", {"codigo": codigo or "(novo)", "nome": nome}

    if entity == "encomendas":
        cliente_nome = _norm_str(row.get("cliente"))
        if not cliente_nome:
            return "error", "Cliente obrigatório"
        cliente = await clientes_repo.find_one({"nome": cliente_nome})
        if not cliente:
            cliente = await clientes_repo.find_one({"codigo": cliente_nome})
        if not cliente:
            # match case-insensitive por nome
            cliente = await clientes_repo.find_one({"nome": {"$regex": f"^{re.escape(cliente_nome)}$", "$options": "i"}})
        if not cliente:
            return "error", f"Cliente não encontrado: {cliente_nome}"

        numero = _norm_code(row.get("numero"))
        estado_raw = _norm_str(row.get("estado") or "aberta").lower().replace(" ", "_")
        estado_map = {
            "aberta": "aberta",
            "em_producao": "em_producao",
            "emprodução": "em_producao",
            "em_produção": "em_producao",
            "concluida": "concluida",
            "concluída": "concluida",
            "cancelada": "cancelada",
            "open": "aberta",
            "closed": "concluida",
            "cancelled": "cancelada",
            "canceled": "cancelada",
        }
        estado = estado_map.get(estado_raw, "aberta")

        orc_num = _norm_code(row.get("orcamento_numero"))
        orcamento_id = None
        if orc_num:
            orc = await orcamentos_repo.find_one({"numero": orc_num})
            if orc:
                orcamento_id = orc.get("id")
            else:
                return "error", f"Orçamento não encontrado: {orc_num}"

        valor_total = row.get("valor_total")
        has_valor = valor_total is not None and _norm_str(valor_total) != ""
        patch = {
            "cliente": cliente.get("nome") or cliente_nome,
            "cliente_id": cliente.get("id"),
            "descricao": _norm_str(row.get("descricao")),
            "data": _parse_date(row.get("data")),
            "prazo_entrega": _parse_date(row.get("prazo_entrega")),
            "estado": estado,
            "notas": _norm_str(row.get("notas")),
            "valor_pago": _parse_float(row.get("valor_pago")),
            "desconto_total": _parse_float(row.get("desconto_total")),
            "desconto_total_tipo": (
                "eur" if _norm_str(row.get("desconto_total_tipo") or "pct").lower() in ("eur", "€", "euro") else "pct"
            ),
            "envio": _parse_float(row.get("envio")),
            "orcamento_numero": orc_num or None,
            "orcamento_id": orcamento_id,
        }
        if has_valor:
            patch["valor_total"] = _parse_float(valor_total)
            patch["valor_total_manual"] = True

        existing = cache.get(numero) if numero else None
        if existing is None and numero:
            existing = await encomendas_repo.find_one({"numero": numero})
        if existing and not existing.get("_pending"):
            if not allow_update:
                return "skip", {
                    "codigo": existing.get("numero") or numero,
                    "nome": existing.get("cliente") or patch["cliente"],
                    "motivo": "já existe — use modo Actualizar ou coluna atualizar=sim",
                }
            if not dry_run:
                await encomendas_repo.update(existing["id"], patch)
            return "update", {"codigo": existing.get("numero") or numero, "nome": patch["cliente"]}
        if not dry_run:
            doc = {
                "id": new_id(),
                "numero": numero or await next_sequence("ENC"),
                **patch,
                "artigos": [],
                "imagens": [],
                "anexos": [],
                "pagamentos": [],
                "autorizada_producao": False,
                "entregue": False,
                "created_at": now_iso(),
            }
            # Evitar sobrescrever desconto/envio com defaults errados se já estavam no patch
            if "desconto_total" not in patch:
                doc["desconto_total"] = 0.0
            if "desconto_total_tipo" not in patch:
                doc["desconto_total_tipo"] = "pct"
            if "envio" not in patch:
                doc["envio"] = 0.0
            if not has_valor:
                doc["valor_total"] = None
                doc["valor_total_manual"] = False
            await encomendas_repo.insert(doc)
            return "create", {"codigo": doc["numero"], "nome": patch["cliente"]}
        return "create", {"codigo": numero or "(novo)", "nome": patch["cliente"]}

    if entity == "materiais":
        if not nome:
            return "error", "Nome obrigatório"
        patch = {
            "nome": nome,
            "unidade": _norm_str(row.get("unidade")) or "un",
            "custo_unitario": _parse_float(row.get("custo_unitario")),
        }
        existing = await consumiveis_repo.find_one({"codigo": codigo}) if codigo else None
        if existing:
            if not allow_update:
                return "skip", {"codigo": existing.get("codigo") or codigo, "nome": existing.get("nome") or nome, "motivo": "já existe"}
            if not dry_run:
                await consumiveis_repo.update(existing["id"], patch)
            return "update", {"codigo": existing.get("codigo") or codigo, "nome": nome}
        if not dry_run:
            doc = {
                "id": new_id(),
                "codigo": codigo or await next_codigo("material"),
                **patch,
                "created_at": now_iso(),
            }
            await consumiveis_repo.insert(doc)
            return "create", {"codigo": doc["codigo"], "nome": nome}
        return "create", {"codigo": codigo or "(novo)", "nome": nome}

    if entity == "maquinas":
        if not nome:
            return "error", "Nome obrigatório"
        patch = {
            "nome": nome,
            "custo_amortizacao_hora": _parse_float(row.get("custo_amortizacao_hora")),
            "custo_energia_hora": _parse_float(row.get("custo_energia_hora")),
        }
        existing = await maquinas_repo.find_one({"codigo": codigo}) if codigo else None
        if existing:
            if not allow_update:
                return "skip", {"codigo": existing.get("codigo") or codigo, "nome": existing.get("nome") or nome, "motivo": "já existe"}
            if not dry_run:
                await maquinas_repo.update(existing["id"], patch)
            return "update", {"codigo": existing.get("codigo") or codigo, "nome": nome}
        if not dry_run:
            doc = {
                "id": new_id(),
                "codigo": codigo or await next_codigo("maquina"),
                **patch,
                "created_at": now_iso(),
            }
            await maquinas_repo.insert(doc)
            return "create", {"codigo": doc["codigo"], "nome": nome}
        return "create", {"codigo": codigo or "(novo)", "nome": nome}

    if entity == "mao_obra":
        if not nome:
            return "error", "Nome obrigatório"
        patch = {
            "nome": nome,
            "custo_hora": _parse_float(row.get("custo_hora")),
            "responsavel_personalizacoes": _parse_bool(row.get("responsavel_personalizacoes")),
        }
        existing = await mao_obra_repo.find_one({"codigo": codigo}) if codigo else None
        if existing:
            if not allow_update:
                return "skip", {"codigo": existing.get("codigo") or codigo, "nome": existing.get("nome") or nome, "motivo": "já existe"}
            if not dry_run:
                await mao_obra_repo.update(existing["id"], patch)
            return "update", {"codigo": existing.get("codigo") or codigo, "nome": nome}
        if not dry_run:
            doc = {
                "id": new_id(),
                "codigo": codigo or await next_codigo("mao_obra"),
                **patch,
                "created_at": now_iso(),
            }
            await mao_obra_repo.insert(doc)
            return "create", {"codigo": doc["codigo"], "nome": nome}
        return "create", {"codigo": codigo or "(novo)", "nome": nome}

    if entity == "categorias":
        if not nome:
            return "error", "Nome obrigatório"
        existing = await categorias_repo.find_one({"codigo": codigo}) if codigo else None
        if not existing:
            existing = await categorias_repo.find_one({"nome": nome})
        if existing:
            if not allow_update:
                return "skip", {"codigo": existing.get("codigo") or codigo, "nome": existing.get("nome") or nome, "motivo": "já existe"}
            if not dry_run:
                await categorias_repo.update(existing["id"], {"nome": nome})
            return "update", {"codigo": existing.get("codigo") or codigo, "nome": nome}
        if not dry_run:
            doc = {
                "id": new_id(),
                "codigo": codigo or await next_codigo("categoria"),
                "nome": nome,
                "created_at": now_iso(),
            }
            await categorias_repo.insert(doc)
            return "create", {"codigo": doc["codigo"], "nome": nome}
        return "create", {"codigo": codigo or "(novo)", "nome": nome}

    if entity == "subcategorias":
        if not nome:
            return "error", "Nome obrigatório"
        cat_ref = _norm_str(row.get("categoria_codigo") or row.get("categoria"))
        cat = await _resolve_categoria(cat_ref)
        if not cat:
            return "error", f"Categoria não encontrada: {cat_ref or '(vazia)'}"
        existing = await subcategorias_repo.find_one({"codigo": codigo}) if codigo else None
        if not existing:
            existing = await subcategorias_repo.find_one({"nome": nome, "categoria_id": cat["id"]})
        patch = {
            "nome": nome,
            "categoria_id": cat["id"],
            "categoria_nome": cat.get("nome") or "",
        }
        if existing:
            if not allow_update:
                return "skip", {"codigo": existing.get("codigo") or codigo, "nome": existing.get("nome") or nome, "motivo": "já existe"}
            if not dry_run:
                await subcategorias_repo.update(existing["id"], patch)
            return "update", {"codigo": existing.get("codigo") or codigo, "nome": nome}
        if not dry_run:
            doc = {
                "id": new_id(),
                "codigo": codigo or await next_codigo("subcategoria"),
                **patch,
                "created_at": now_iso(),
            }
            await subcategorias_repo.insert(doc)
            return "create", {"codigo": doc["codigo"], "nome": nome}
        return "create", {"codigo": codigo or "(novo)", "nome": nome}

    if entity == "artigos":
        if not nome:
            return "error", "Nome obrigatório"
        cat = await _resolve_categoria(_norm_str(row.get("categoria")))
        sub = None
        if cat:
            sub = await _resolve_subcategoria(_norm_str(row.get("subcategoria")), cat["id"])
        forn_ref = _norm_str(row.get("fornecedor"))
        forn = None
        if forn_ref:
            forn = await fornecedores_repo.find_one({"codigo": forn_ref})
            if not forn:
                forn = await fornecedores_repo.find_one({"nome": forn_ref})

        tipo_raw = _norm_str(row.get("tipo_artigo") or "ativo").lower().replace(" ", "_")
        if tipo_raw in ("produzido",):
            tipo_raw = "ativo"
            produzido = True
        else:
            produzido = _parse_bool(row.get("produzido")) if "produzido" in row else False
        if tipo_raw not in ("ativo", "consumivel", "servico", "nao_utilizado", "inativo"):
            return "error", f"tipo_artigo inválido: {tipo_raw}"
        if tipo_raw != "ativo":
            produzido = False

        ativo = _parse_bool(row.get("ativo")) if row.get("ativo") not in (None, "") else (tipo_raw != "inativo")
        if tipo_raw == "inativo":
            ativo = False

        patch = {
            "nome": nome,
            "descricao": _norm_str(row.get("descricao")),
            "unidade": _norm_str(row.get("unidade")) or "un",
            "tipo_artigo": tipo_raw,
            "produzido": produzido,
            "categoria_id": cat["id"] if cat else None,
            "categoria_nome": (cat or {}).get("nome") or "",
            "subcategoria_id": sub["id"] if sub else None,
            "subcategoria_nome": (sub or {}).get("nome") or "",
            "custo_artigo": _parse_float(row.get("custo_artigo")),
            "margem": _parse_float(row.get("margem"), 0.0 if tipo_raw == "consumivel" else 30.0),
            "comissao_pct": _parse_float(row.get("comissao_pct")),
            "fabricante": _norm_str(row.get("fabricante")),
            "cod_fabricante": _norm_code(row.get("cod_fabricante")),
            "fornecedor_id": forn["id"] if forn else None,
            "fornecedor_nome": (forn or {}).get("nome") or forn_ref,
            "cod_fornecedor": _norm_code(row.get("cod_fornecedor")),
            "website": _norm_str(row.get("website")),
            "ficha_produto": _norm_str(row.get("ficha_produto")),
            "plano_contas": _norm_code(row.get("plano_contas")),
            "codigo_produto": _norm_code(row.get("codigo_produto")),
            "comprimento_mm": _parse_float(row.get("comprimento_mm")),
            "largura_mm": _parse_float(row.get("largura_mm")),
            "espessura_mm": _parse_float(row.get("espessura_mm")),
            "peso_kg": _parse_float(row.get("peso_kg")),
            "comprimento_unidade": _norm_str(row.get("comprimento_unidade")) or "mm",
            "largura_unidade": _norm_str(row.get("largura_unidade")) or "mm",
            "espessura_unidade": _norm_str(row.get("espessura_unidade")) or "mm",
            "peso_unidade": _norm_str(row.get("peso_unidade")) or "kg",
            "responsavel": _norm_str(row.get("responsavel")),
            "qtd_uni": _parse_float(row.get("qtd_uni"), 1.0),
            "qtd_stock": _parse_float(row.get("qtd_stock")),
            "nivel_reabastecimento": _parse_float(row.get("nivel_reabastecimento")),
            "qtd_ultima_compra": _parse_float(row.get("qtd_ultima_compra")),
            "diversos": _parse_bool(row.get("diversos")),
            "ativo": ativo,
        }
        existing = await artigos_repo.find_one({"codigo": codigo}) if codigo else None
        if existing:
            if not allow_update:
                return "skip", {"codigo": existing.get("codigo") or codigo, "nome": existing.get("nome") or nome, "motivo": "já existe"}
            if not dry_run:
                await artigos_repo.update(existing["id"], patch)
            return "update", {"codigo": existing.get("codigo") or codigo, "nome": nome}
        if not dry_run:
            doc = {
                "id": new_id(),
                "codigo": codigo or await next_codigo("artigo"),
                **patch,
                "imagem": "",
                "materiais": [],
                "roteiro": [],
                "created_at": now_iso(),
                "updated_at": now_iso(),
            }
            await artigos_repo.insert(doc)
            return "create", {"codigo": doc["codigo"], "nome": nome}
        return "create", {"codigo": codigo or "(novo)", "nome": nome}

    return "error", f"Entidade desconhecida: {entity}"


async def _parse_encomenda_linha(row: Dict[str, Any]):
    """Devolve dict da linha ou string de erro."""
    cod = _norm_code(row.get("artigo_codigo"))
    nome = _norm_str(row.get("artigo_nome"))
    art = None
    if cod:
        art = await artigos_repo.find_one({"codigo": cod})
    if not art and nome:
        art = await artigos_repo.find_one({"nome": nome})
        if not art:
            art = await artigos_repo.find_one({"nome": {"$regex": f"^{re.escape(nome)}$", "$options": "i"}})
    if not nome and art:
        nome = art.get("nome") or ""
    if not nome and not art:
        return "Indique artigo_codigo ou artigo_nome"
    qtd = _parse_float(row.get("quantidade"), 1.0)
    if qtd <= 0:
        return "Quantidade deve ser > 0"
    desc_tipo = _norm_str(row.get("desconto_tipo") or "pct").lower()
    if desc_tipo not in ("pct", "eur"):
        desc_tipo = "pct"
    return {
        "id": new_id(),
        "artigo_id": art["id"] if art else None,
        "artigo_nome": nome or (art or {}).get("nome") or "",
        "imagem": (art or {}).get("imagem") or "",
        "quantidade": qtd,
        "preco_unit": _parse_float(row.get("preco_unit"), _parse_float((art or {}).get("custo_artigo"), 0.0)),
        "desconto": _parse_float(row.get("desconto")),
        "desconto_tipo": desc_tipo,
        "personalizacoes": _pers_from_cell(row.get("personalizacoes")),
    }


async def _apply_encomenda_linhas(
    numero: str,
    linhas: List[dict],
    *,
    dry_run: bool,
    mode: str = "create",
    known_new: bool = False,
) -> Tuple[str, Any]:
    """Liga linhas a uma encomenda. Em create: só preenche se ainda não tiver artigos."""
    numero = _norm_code(numero)
    enc = await encomendas_repo.find_one({"numero": numero})
    if not enc:
        if dry_run or known_new:
            return "create", {"codigo": numero, "nome": f"{len(linhas)} linha(s)"}
        return "error", f"Encomenda não encontrada: {numero}"

    existing_arts = enc.get("artigos") or []
    allow_update = mode == "update"
    if existing_arts and not allow_update and not known_new:
        return "skip", {
            "codigo": numero,
            "nome": f"{len(existing_arts)} linha(s) já existentes",
            "motivo": "encomenda já tem linhas — use modo Actualizar para substituir",
        }

    if not dry_run:
        await encomendas_repo.update(enc["id"], {"artigos": linhas})
    acao = "update" if existing_arts else "create"
    return acao, {
        "codigo": numero,
        "nome": f"{len(linhas)} linha(s)",
    }
